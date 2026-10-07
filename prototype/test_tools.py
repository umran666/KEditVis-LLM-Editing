"""Regressions for offline analysis and document export helpers."""
import math
from pathlib import Path
import tempfile
import unittest

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from analyze_schemes import scheme_activity_score, scheme_projection_score, spearman_rank_correlation, target_token_probs
from error_analysis import summarize_fact
from export_doc import build_revised_document, REVISED_ABSTRACT_TEXT


class AnalysisTests(unittest.TestCase):
    def test_incomplete_or_nonfinite_activity_is_rejected(self):
        signals = [{"layer": 2, "cosine_similarity": .2}]
        with self.assertRaises(ValueError):
            scheme_activity_score(signals, [2, 3])
        with self.assertRaises(ValueError):
            scheme_activity_score(signals * 2, [2])
        with self.assertRaises(ValueError):
            scheme_activity_score([{"layer": 2, "cosine_similarity": math.nan}], [2])

    def test_projection_deduplicates_layers_and_keeps_missing_last_data_missing(self):
        signals = [{"layer": 2, "last_top_tokens": [{"token": " Rome", "prob": .8}]}]
        self.assertEqual(scheme_projection_score(signals, [2, 2], "Rome")["span_overlap"], 1)
        missing = [{"layer": 2, "last_top_tokens": [], "top_tokens": [{"token": "Rome", "prob": .8}]}]
        self.assertEqual(target_token_probs(missing, "Rome"), {})

    def test_correlations_reject_nonfinite_pairs(self):
        for value in [math.nan, math.inf, -math.inf]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                spearman_rank_correlation([1, value, 3], [1, 2, 3])

    def test_cosine_validation_accepts_float32_roundoff(self):
        signal = {"layer": 0, "cosine_similarity": 1.0000003576278687,
                  "last_top_tokens": []}
        self.assertGreater(scheme_activity_score([signal], [0])["mean_abs_cos_sim"], 1)
        self.assertEqual(summarize_fact({"fact": {"target_new": "Rome"},
                                       "baseline": {"layer_signals": [signal]}})["n_layers"], 1)
        with self.assertRaises(ValueError):
            scheme_activity_score([{**signal, "cosine_similarity": 1.01}], [0])

    def test_error_summary_uses_last_token_projection(self):
        fact = {"fact": {"target_new": "Rome"}, "baseline": {"layer_signals": [
            {"layer": 0, "cosine_similarity": .3, "top_tokens": [{"token": "Tower", "prob": .5}],
             "last_top_tokens": [{"token": " Rome", "prob": .6}]}]}}
        result = summarize_fact(fact)
        self.assertEqual(result["layers_with_target_in_top5"], [0])
        self.assertEqual(result["n_layers"], 1)


class DocumentExportTests(unittest.TestCase):
    def test_export_preserves_source_and_inherited_or_left_alignment(self):
        with tempfile.TemporaryDirectory() as folder:
            for alignment in [None, WD_ALIGN_PARAGRAPH.LEFT]:
                with self.subTest(alignment=alignment):
                    source = Path(folder) / "source.docx"
                    output = Path(folder) / "revised.docx"
                    document = docx.Document()
                    paragraph = document.add_paragraph()
                    paragraph.alignment = alignment
                    run = paragraph.add_run("Large Language Models frequently encode obsolete facts.")
                    run.font.name = "Arial"
                    run.font.size = Pt(11)
                    run.italic = True
                    document.save(source)
                    original = source.read_bytes()
                    build_revised_document(source, output)
                    self.assertEqual(source.read_bytes(), original)
                    revised = docx.Document(output).paragraphs[0]
                    self.assertEqual(revised.text, REVISED_ABSTRACT_TEXT)
                    self.assertEqual(revised.alignment, alignment)
                    self.assertEqual(revised.runs[0].font.name, "Arial")
                    self.assertEqual(revised.runs[0].font.size, Pt(11))
                    self.assertTrue(revised.runs[0].italic)
                    with self.assertRaises(ValueError):
                        build_revised_document(source, source)
                    self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
