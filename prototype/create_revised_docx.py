"""
Generates Knowledge_Editing_LLMs_Revised.docx from Knowledge_Editing_LLMs.docx.
Preserves all original student names, registration numbers, batch numbers, guide details,
margins, fonts, alignments, and base paper citation while updating the abstract paragraph
to reconcile scientific claims with empirical benchmark findings.
The original Knowledge_Editing_LLMs.docx is left strictly untouched.
"""

from pathlib import Path
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

ORIGINAL_PATH = Path(r"c:\Users\shaik\Research\LLM Editing\Knowledge_Editing_LLMs.docx")
REVISED_PATH = Path(r"c:\Users\shaik\Research\LLM Editing\Knowledge_Editing_LLMs_Revised.docx")

REVISED_ABSTRACT_TEXT = (
    "Large Language Models frequently encode obsolete or inaccurate factual associations "
    "within their parameters. While locate-then-edit methods such as ROME and MEMIT provide "
    "computationally efficient alternatives to full model retraining, conventional pipelines "
    "rely on static, model-wide layer presets. These fixed presets ignore fact-specific "
    "activation patterns, frequently causing incomplete edits, localized hallucination, or "
    "catastrophic parameter drift. Grounded in recent visual analytics research for model editing, "
    "this capstone project develops an interactive prototype for human-in-the-loop layer "
    "selection. The system extracts internal model signals, specifically layer-wise residual "
    "variance and vocabulary probability distributions, presenting them through coordinated visual "
    "interfaces. Users can evaluate candidate layer ranges across editing success, paraphrase "
    "generalization, and neighborhood locality metrics. The framework incorporates a reversible "
    "model state mechanism and dimensionality reduction to monitor hidden state drift. "
    "Experiments conducted on open-source Transformer architectures using standardized editing "
    "benchmarks demonstrate that interactive telemetry-guided layer selection successfully identifies "
    "viable editing bands to prevent catastrophic failure modes, while multi-context optimization "
    "enhances paraphrase generalization under bounded parameter drift."
)


def main():
    if not ORIGINAL_PATH.exists():
        raise FileNotFoundError(f"Original file not found: {ORIGINAL_PATH}")

    doc = docx.Document(ORIGINAL_PATH)
    
    # Locate paragraph 32 (the ABSTRACT body paragraph)
    # Verify it matches the original abstract start
    target_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Large Language Models frequently encode obsolete"):
            target_idx = i
            break
            
    if target_idx is None:
        raise ValueError("Could not locate original abstract body paragraph.")
        
    p = doc.paragraphs[target_idx]
    
    # Preserve formatting properties
    align = p.alignment or WD_ALIGN_PARAGRAPH.JUSTIFY
    orig_runs = p.runs
    font_name = orig_runs[0].font.name if orig_runs and orig_runs[0].font.name else "Times New Roman"
    font_size = orig_runs[0].font.size if orig_runs and orig_runs[0].font.size else docx.shared.Pt(12)
    
    # Replace text
    p.text = REVISED_ABSTRACT_TEXT
    p.alignment = align
    for run in p.runs:
        run.font.name = font_name
        run.font.size = font_size
        
    doc.save(REVISED_PATH)
    print(f"Successfully generated revised abstract document at:\n{REVISED_PATH}")


if __name__ == "__main__":
    main()
